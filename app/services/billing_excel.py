"""Excel de cobranza: exportar los pagos fuera de contratos para analizarlos y volver a importarlos ya ordenados.

Flujo (PRP 008):
1. `export_workbook()` arma un .xlsx con todos los pagos que no están aplicados a cuotas (históricos sin clasificar
   y pagos sueltos), una propuesta de contrato por paciente con plan heredado o pagos huérfanos, los contratos
   activos (para poder apuntar a ellos), una hoja de resumen con fórmulas y listas para los desplegables.
2. El admin decide en las columnas amarillas qué es cada pago (cuotas de un contrato, pago suelto o dejarlo) y
   ajusta los contratos propuestos (monto, cuotas, día de pago, primera cuota).
3. `import_workbook(..., apply=False)` simula todo y devuelve fila por fila qué pasaría (sin tocar la base).
   Con `apply=True` hace lo mismo de verdad en una sola transacción y deja un `BillingImport` con la huella del
   archivo; cada contrato, pago y asignación creados guardan el lote. Si hay errores, no se aplica nada.

Re-importar el mismo archivo no duplica: un pago ya aplicado se salta y un pago nuevo idéntico (paciente, fecha,
monto y referencia) se marca como posible duplicado.
"""

import hashlib
import io
import json
import statistics
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.extensions import db
from app.models import Contract, Payment, Sede, User
from app.models.billing import BillingImport, PaymentAllocation
from app.services import billing

DESTINOS = ('CUOTAS', 'SUELTO', 'DEJAR')
FRECUENCIAS = ('Mensual', 'Quincenal')
METODOS = ('Yape', 'Plin', 'Efectivo', 'Transferencia', 'Tarjeta', 'Otro')
CATEGORIAS = tuple(billing.ONE_OFF_CATEGORIES)

PAY_COLS = [
    # (encabezado, ancho, editable)
    ('ID pago', 9, False),
    ('Fecha', 12, False),
    ('ID paciente', 11, False),
    ('Código', 11, False),
    ('Paciente', 28, False),
    ('Sede', 14, False),
    ('Monto', 11, False),
    ('Descuento', 10, False),
    ('Método', 13, False),
    ('Referencia', 16, False),
    ('Notas', 28, False),
    ('Estado actual', 15, False),
    ('Destino', 11, True),
    ('Categoría (si es suelto)', 20, True),
    ('Contrato (clave)', 16, True),
    ('Cuota n° (opcional)', 12, True),
    ('Comentario', 28, True),
]
CON_COLS = [
    ('Clave', 14, True),
    ('ID paciente', 11, False),
    ('Código', 11, False),
    ('Paciente', 28, False),
    ('Sede', 14, False),
    ('Estado', 11, False),
    ('Nombre del plan', 26, True),
    ('Monto de cuota', 13, True),
    ('N° de cuotas (vacío = renovable)', 16, True),
    ('Frecuencia', 12, True),
    ('Día de pago', 10, True),
    ('Primera cuota', 13, True),
    ('Inicio del servicio', 13, True),
    ('Notas', 30, True),
    ('Plan actual (referencia)', 30, False),
    ('Pagos asignados en este archivo', 16, False),
]

GREEN = PatternFill('solid', fgColor='3F6212')
GREY = PatternFill('solid', fgColor='F1F3EC')
YELLOW = PatternFill('solid', fgColor='FFF4C2')
THIN = Side(style='thin', color='D3D8C8')
MONEY = '"S/" #,##0.00'
DATE = 'DD/MM/YYYY'


def _key(patient):
    return f'C-{patient.login_code or patient.id}'


def _plan_text(u):
    if not (u.payment_amount or 0) > 0:
        return ''
    freq = 'quincenal' if u.payment_plan == 'quincenal' else 'mensual'
    due = u.payment_due_date.strftime('%d/%m/%Y') if u.payment_due_date else 'sin fecha'
    return f'S/ {u.payment_amount:.2f} {freq} · vence {due}'


def _header(ws, cols):
    for i, (title, width, editable) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = Font(bold=True, color='FFFFFF' if not editable else '3B2F00')
        c.fill = GREEN if not editable else PatternFill('solid', fgColor='F5C842')
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}1'


def _paint(ws, cols, row):
    for i, (_, _, editable) in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i)
        cell.fill = YELLOW if editable else GREY
        cell.border = Border(bottom=THIN)


def export_workbook():
    """Bytes del .xlsx con los pagos fuera de contratos y las propuestas de contrato."""
    sedes = dict(db.session.query(Sede.id, Sede.name))
    patients = {u.id: u for u in User.query.filter(User.role == 'jugador').order_by(User.username)}
    active = defaultdict(list)
    for c in Contract.query.filter(Contract.status == 'active', Contract.is_active.isnot(False)).order_by(Contract.id):
        active[c.patient_id].append(c)
    loose = (
        Payment.query.filter(
            Payment.is_active.isnot(False),
            db.or_(Payment.kind.in_(('unassigned', 'one_off')), Payment.kind.is_(None)),
        )
        .order_by(Payment.date, Payment.id)
        .all()
    )
    loose = [p for p in loose if billing.allocated(p.id) <= billing.EPS]
    by_patient = defaultdict(list)
    for p in loose:
        by_patient[p.patient_id].append(p)

    wb = Workbook()
    ins = wb.active
    ins.title = 'Instrucciones'
    pay = wb.create_sheet('Pagos')
    con = wb.create_sheet('Contratos')
    pac = wb.create_sheet('Pacientes')
    res = wb.create_sheet('Resumen')
    lst = wb.create_sheet('Listas')

    # ── Listas (desplegables) ──
    for col, (title, values) in enumerate(
        (('Destino', DESTINOS), ('Categoría', CATEGORIAS), ('Frecuencia', FRECUENCIAS), ('Método', METODOS)), start=1
    ):
        lst.cell(row=1, column=col, value=title).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            lst.cell(row=r, column=col, value=v)
    lst.sheet_state = 'hidden'

    # ── Contratos: propuestas + existentes ──
    _header(con, CON_COLS)
    row = 2
    proposed = []
    for pid, u in patients.items():
        if active.get(pid):
            continue
        mine = by_patient.get(pid, [])
        has_plan = (u.payment_amount or 0) > 0
        if not has_plan and not mine:
            continue
        amounts = [round(float(p.amount or 0), 2) for p in mine if p.kind != 'one_off']
        amount = (
            round(float(u.payment_amount), 2)
            if has_plan
            else (
                Counter(amounts).most_common(1)[0][0]
                if amounts
                else round(statistics.median([p.amount for p in mine]), 2)
            )
        )
        first_pay = mine[0].date.date() if mine else None
        day = u.payment_day or (u.payment_due_date.day if u.payment_due_date else (first_pay.day if first_pay else 1))
        anchor = first_pay or u.payment_due_date or date.today()
        first_due = date(anchor.year, anchor.month, min(day, 28))
        proposed.append((u, amount, day, first_due, has_plan))
    for u, amount, day, first_due, has_plan in proposed:
        values = [
            _key(u),
            u.id,
            u.login_code or '',
            u.username or u.email,
            sedes.get(u.sede_id, ''),
            'NUEVO',
            'Plan mensual',
            amount,
            None,
            'Quincenal' if u.payment_plan == 'quincenal' else 'Mensual',
            day,
            first_due,
            first_due,
            '' if has_plan else 'Revisar: propuesto desde sus pagos, no tenía plan registrado.',
            _plan_text(u),
            f'=SUMIFS(Pagos!G:G,Pagos!O:O,A{row},Pagos!M:M,"CUOTAS")',
        ]
        for col, v in enumerate(values, start=1):
            con.cell(row=row, column=col, value=v)
        _paint(con, CON_COLS, row)
        row += 1
    for pid, contracts in active.items():
        u = patients.get(pid)
        if not u:
            continue
        for c in contracts:
            values = [
                f'#{c.id}',
                u.id,
                u.login_code or '',
                u.username or u.email,
                sedes.get(u.sede_id, ''),
                'EXISTENTE',
                c.name,
                float(c.installment_amount or 0),
                None if c.renewable else c.installment_count,
                c.billing_type or 'Mensual',
                c.payment_day,
                c.payment_start_date,
                c.service_start_date,
                'Contrato ya registrado: solo sirve para apuntar pagos (no se modifica).',
                _plan_text(u),
                f'=SUMIFS(Pagos!G:G,Pagos!O:O,A{row},Pagos!M:M,"CUOTAS")',
            ]
            for col, v in enumerate(values, start=1):
                con.cell(row=row, column=col, value=v)
            _paint(con, CON_COLS, row)
            for col in range(1, len(CON_COLS) + 1):
                con.cell(row=row, column=col).fill = GREY
            row += 1
    last_con = max(row - 1, 2)
    for r in range(2, last_con + 1):
        con[f'H{r}'].number_format = MONEY
        con[f'L{r}'].number_format = DATE
        con[f'M{r}'].number_format = DATE
        con[f'P{r}'].number_format = MONEY
    dv = DataValidation(type='list', formula1='=Listas!$C$2:$C$3', allow_blank=True)
    con.add_data_validation(dv)
    dv.add(f'J2:J{last_con + 200}')
    dv_day = DataValidation(type='whole', operator='between', formula1='1', formula2='31', allow_blank=True)
    dv_day.error = 'El día de pago va de 1 a 31.'
    con.add_data_validation(dv_day)
    dv_day.add(f'K2:K{last_con + 200}')

    # ── Pagos ──
    _header(pay, PAY_COLS)
    suggested_key = {u.id: _key(u) for u, *_ in proposed}
    for pid, contracts in active.items():
        if len(contracts) == 1:
            suggested_key[pid] = f'#{contracts[0].id}'
    row = 2
    for p in loose:
        u = patients.get(p.patient_id)
        state = 'Pago suelto' if p.kind == 'one_off' else 'Sin clasificar'
        dest = 'SUELTO' if p.kind == 'one_off' else ('CUOTAS' if p.patient_id in suggested_key else '')
        values = [
            p.id,
            p.date.date() if p.date else None,
            p.patient_id,
            (u.login_code if u else '') or '',
            (u.username or u.email) if u else f'Paciente #{p.patient_id}',
            sedes.get(u.sede_id, '') if u else '',
            float(p.amount or 0),
            float(p.discount or 0),
            p.method,
            p.reference or '',
            p.notes or '',
            state,
            dest,
            p.category or '',
            suggested_key.get(p.patient_id, '') if dest == 'CUOTAS' else '',
            None,
            '',
        ]
        for col, v in enumerate(values, start=1):
            pay.cell(row=row, column=col, value=v)
        _paint(pay, PAY_COLS, row)
        row += 1
    last_pay = max(row - 1, 2)
    for r in range(2, last_pay + 300):
        pay[f'B{r}'].number_format = DATE
        pay[f'G{r}'].number_format = MONEY
        pay[f'H{r}'].number_format = MONEY
    for col, ref in (('M', '$A$2:$A$4'), ('N', '$B$2:$B$5'), ('I', '$D$2:$D$7')):
        v = DataValidation(type='list', formula1=f'=Listas!{ref}', allow_blank=True)
        pay.add_data_validation(v)
        v.add(f'{col}2:{col}{last_pay + 300}')
    keys = DataValidation(type='list', formula1=f'=Contratos!$A$2:$A${last_con + 200}', allow_blank=True)
    keys.error = 'Usa una clave de la hoja Contratos.'
    pay.add_data_validation(keys)
    keys.add(f'O2:O{last_pay + 300}')

    # ── Pacientes (referencia) ──
    pcols = [
        ('ID paciente', 11, False),
        ('Código', 11, False),
        ('Paciente', 30, False),
        ('Sede', 16, False),
        ('Estado', 12, False),
        ('Plan actual', 32, False),
        ('Contrato activo', 16, False),
    ]
    _header(pac, pcols)
    for r, u in enumerate(patients.values(), start=2):
        for col, v in enumerate(
            [
                u.id,
                u.login_code or '',
                u.username or u.email,
                sedes.get(u.sede_id, ''),
                u.account_status or 'active',
                _plan_text(u),
                ', '.join(f'#{c.id}' for c in active.get(u.id, [])),
            ],
            start=1,
        ):
            pac.cell(row=r, column=col, value=v)

    # ── Resumen (fórmulas vivas: cambian al editar el archivo) ──
    res.column_dimensions['A'].width = 34
    res.column_dimensions['B'].width = 16
    res.column_dimensions['C'].width = 12
    res['A1'] = 'Resumen de pagos fuera de contratos'
    res['A1'].font = Font(bold=True, size=14)
    res['A2'] = f'Exportado el {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    rows = [
        ('Total de pagos en el archivo', '=SUM(Pagos!G:G)', '=COUNT(Pagos!G:G)'),
        ('Irán a cuotas (CUOTAS)', '=SUMIFS(Pagos!G:G,Pagos!M:M,"CUOTAS")', '=COUNTIFS(Pagos!M:M,"CUOTAS")'),
        ('Pagos sueltos (SUELTO)', '=SUMIFS(Pagos!G:G,Pagos!M:M,"SUELTO")', '=COUNTIFS(Pagos!M:M,"SUELTO")'),
        ('Se dejan como están (DEJAR)', '=SUMIFS(Pagos!G:G,Pagos!M:M,"DEJAR")', '=COUNTIFS(Pagos!M:M,"DEJAR")'),
        ('Sin decidir (Destino vacío)', '=SUM(Pagos!G:G)-B5-B6-B7', '=C4-C5-C6-C7'),
        ('Contratos nuevos propuestos', '', '=COUNTIFS(Contratos!F:F,"NUEVO")'),
    ]
    res['A3'], res['B3'], res['C3'] = 'Concepto', 'Monto', 'Pagos'
    for c in ('A3', 'B3', 'C3'):
        res[c].font = Font(bold=True)
    for i, (label, money, count) in enumerate(rows, start=4):
        res[f'A{i}'], res[f'B{i}'], res[f'C{i}'] = label, money, count
        res[f'B{i}'].number_format = MONEY
    res['A12'] = 'Por mes (fecha del pago)'
    res['A12'].font = Font(bold=True)
    months = sorted({(p.date.year, p.date.month) for p in loose if p.date})
    for i, (y, m) in enumerate(months, start=13):
        start = date(y, m, 1)
        end = date(y + (m == 12), m % 12 + 1, 1)
        res[f'A{i}'] = start
        res[f'A{i}'].number_format = 'MMM YYYY'
        res[f'B{i}'] = (
            f'=SUMIFS(Pagos!G:G,Pagos!B:B,">="&DATE({start.year},{start.month},1),Pagos!B:B,"<"&DATE({end.year},{end.month},1))'
        )
        res[f'B{i}'].number_format = MONEY
        res[f'C{i}'] = (
            f'=COUNTIFS(Pagos!B:B,">="&DATE({start.year},{start.month},1),Pagos!B:B,"<"&DATE({end.year},{end.month},1))'
        )
    base = 13 + len(months) + 1
    res[f'A{base}'] = 'Por método'
    res[f'A{base}'].font = Font(bold=True)
    for i, method in enumerate(sorted({p.method for p in loose if p.method}), start=base + 1):
        res[f'A{i}'] = method
        res[f'B{i}'] = f'=SUMIFS(Pagos!G:G,Pagos!I:I,A{i})'
        res[f'B{i}'].number_format = MONEY
        res[f'C{i}'] = f'=COUNTIFS(Pagos!I:I,A{i})'

    # ── Instrucciones ──
    ins.column_dimensions['A'].width = 110
    lines = [
        ('Cobranza: pagos fuera de contratos', True),
        (
            'Este archivo sirve para analizar y para importar. Las columnas AMARILLAS se editan; las grises son datos del sistema (no las cambies).',
            False,
        ),
        ('', False),
        ('1. Hoja «Contratos»', True),
        (
            '· Cada fila NUEVO es un contrato propuesto para un paciente que pagaba sin contrato (o con plan heredado).',
            False,
        ),
        (
            '· Ajusta monto de cuota, n° de cuotas (vacío = renovable: se genera una cuota por mes), frecuencia, día de pago y la fecha de la PRIMERA cuota.',
            False,
        ),
        (
            '· Consejo: pon la primera cuota en el mes del primer pago histórico, así esos pagos cubren las cuotas de su mes y nada queda como deuda falsa.',
            False,
        ),
        (
            '· Si no quieres crear un contrato propuesto, borra su fila. Las filas EXISTENTE no se modifican: solo sirven para apuntar pagos.',
            False,
        ),
        ('', False),
        ('2. Hoja «Pagos»', True),
        (
            '· Destino = CUOTAS: el pago se aplica a las cuotas del contrato indicado en «Contrato (clave)», de la más antigua a la más reciente (o desde la «Cuota n°» si la pones). Lo que sobre queda como saldo a favor.',
            False,
        ),
        (
            '· Destino = SUELTO: no es cuota (evaluación, material, inscripción u otro). Elige la categoría. Cuenta como ingreso, no como cobranza.',
            False,
        ),
        ('· Destino = DEJAR: el pago queda como está, sin clasificar.', False),
        (
            '· Para AGREGAR pagos que no están en el sistema, usa una fila nueva con «ID pago» vacío y llena Fecha, ID paciente (o Código), Monto, Método y Destino.',
            False,
        ),
        ('', False),
        ('3. Importar', True),
        (
            '· En Finanzas → Pagos → «Importar Excel», sube el archivo: primero verás una VISTA PREVIA fila por fila (no cambia nada).',
            False,
        ),
        ('· Si no hay errores, confirma y se aplica todo junto. Si algo falla, no se aplica nada.', False),
        ('· Puedes importar el mismo archivo otra vez: lo ya aplicado se salta, no se duplica.', False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ins.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True)

    wb.active = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── importación ──────────────────────────────────────────────────────────────


def _cell_date(v):
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'Fecha no válida: «{s}» (usa DD/MM/AAAA)')


def _num(v, allow_empty=True):
    if v is None or str(v).strip() == '':
        if allow_empty:
            return None
        raise ValueError('falta el monto')
    return round(float(str(v).replace('S/', '').replace(',', '').strip()), 2)


def _rows(ws, cols):
    titles = [c[0] for c in cols]
    head = [str(c.value or '').strip() for c in ws[1]]
    idx = {t: head.index(t) for t in titles if t in head}
    missing = [t for t in titles if t not in idx]
    if missing:
        raise billing.BillingError(f'A la hoja «{ws.title}» le faltan columnas: {", ".join(missing)}.', 'format')
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, '') for v in row):
            continue
        yield r, {t: row[i] if i < len(row) else None for t, i in idx.items()}


def _patient(rec, by_code):
    pid = rec.get('ID paciente')
    if pid not in (None, ''):
        u = db.session.get(User, int(pid))
        if u and u.role == 'jugador':
            return u
    code = str(rec.get('Código') or '').strip().upper()
    if code and code in by_code:
        return by_code[code]
    raise ValueError('paciente no encontrado (revisa «ID paciente» o «Código»)')


def import_workbook(data, *, file_name='', apply=False, user_id=None):
    """Simula (o aplica) el Excel. Devuelve {'ok', 'applied', 'summary', 'rows'}; nunca aplica si hay errores."""
    sha = hashlib.sha256(data).hexdigest()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise billing.BillingError(
            'No se pudo leer el archivo: ¿es el .xlsx exportado desde el sistema?', 'format'
        ) from exc
    for name in ('Pagos', 'Contratos'):
        if name not in wb.sheetnames:
            raise billing.BillingError(f'Falta la hoja «{name}». Usa el archivo exportado desde el sistema.', 'format')

    batch = None
    if apply:
        batch = BillingImport(file_name=file_name[:255], file_sha256=sha, created_by_id=user_id)
        db.session.add(batch)
        db.session.flush()
    import_id = batch.id if batch else None
    by_code = {u.login_code.upper(): u for u in User.query.filter(User.role == 'jugador', User.login_code.isnot(None))}
    out_rows, errors = [], 0
    summary = Counter()
    money = defaultdict(float)
    contracts = {}
    touched = set()

    def note(sheet, row, status, message):
        nonlocal errors
        out_rows.append({'sheet': sheet, 'row': row, 'status': status, 'message': message})
        if status == 'error':
            errors += 1

    try:
        # 1) Contratos
        for r, rec in _rows(wb['Contratos'], CON_COLS):
            key = str(rec.get('Clave') or '').strip()
            if not key:
                note('Contratos', r, 'error', 'falta la clave')
                continue
            if key.startswith('#'):
                c = db.session.get(Contract, int(key[1:])) if key[1:].isdigit() else None
                if not c or c.status != 'active':
                    note('Contratos', r, 'error', f'el contrato {key} no existe o no está activo')
                    continue
                contracts[key] = c
                continue
            try:
                u = _patient(rec, by_code)
                existing = Contract.query.filter(
                    Contract.patient_id == u.id, Contract.status == 'active', Contract.notes.like(f'%[{key}]%')
                ).first()
                if existing:
                    contracts[key] = existing
                    note(
                        'Contratos',
                        r,
                        'skip',
                        f'{key}: ya se creó en una importación anterior (contrato #{existing.id})',
                    )
                    continue
                n = rec.get('N° de cuotas (vacío = renovable)')
                first_due = _cell_date(rec.get('Primera cuota'))
                if not first_due:
                    raise ValueError('falta la fecha de la primera cuota')
                c = billing.create_plan(
                    u.id,
                    _num(rec.get('Monto de cuota'), allow_empty=False),
                    rec.get('Día de pago') or first_due.day,
                    first_due,
                    installments=int(n) if n not in (None, '') else None,
                    billing_type=str(rec.get('Frecuencia') or 'Mensual').strip().capitalize(),
                    name=(str(rec.get('Nombre del plan') or '').strip() or None),
                    service_start=_cell_date(rec.get('Inicio del servicio')),
                    notes=f'{(rec.get("Notas") or "").strip()} [{key}] (importado desde Excel)'.strip(),
                    import_id=import_id,
                    user_id=user_id,
                    commit=False,
                )
                contracts[key] = c
                touched.add(u.id)
                summary['contracts_created'] += 1
                kind = 'renovable' if c.renewable else f'{c.installment_count} cuotas'
                note(
                    'Contratos',
                    r,
                    'ok',
                    f'{key}: contrato nuevo ({kind}) de S/ {c.installment_amount:.2f} desde {c.payment_start_date:%d/%m/%Y}',
                )
            except (ValueError, billing.BillingError) as e:
                note('Contratos', r, 'error', f'{key}: {e}')

        # 2) Pagos (en orden de fecha: así las cuotas se cubren cronológicamente)
        pay_rows = list(_rows(wb['Pagos'], PAY_COLS))

        def when(item):
            try:
                return _cell_date(item[1].get('Fecha')) or date.max
            except ValueError:
                return date.max

        for r, rec in sorted(pay_rows, key=when):
            dest = str(rec.get('Destino') or '').strip().upper()
            pid = rec.get('ID pago')
            try:
                if dest in ('', 'DEJAR'):
                    summary['left_as_is'] += 1
                    note('Pagos', r, 'skip', 'se deja como está' if dest else 'sin destino: se deja como está')
                    continue
                if dest not in DESTINOS:
                    raise ValueError(f'destino «{dest}» no válido (CUOTAS, SUELTO o DEJAR)')
                key = str(rec.get('Contrato (clave)') or '').strip()
                number = rec.get('Cuota n° (opcional)')
                category = str(rec.get('Categoría (si es suelto)') or '').strip().lower()
                category = {'evaluación': 'evaluacion', 'inscripción': 'inscripcion'}.get(category, category)

                if pid not in (None, ''):
                    p = db.session.get(Payment, int(pid))
                    if not p or p.is_active is False:
                        raise ValueError(f'el pago #{pid} no existe')
                    if billing.allocated(p.id) > billing.EPS:
                        summary['already_applied'] += 1
                        note('Pagos', r, 'skip', f'pago #{p.id}: ya estaba aplicado a cuotas')
                        continue
                    if dest == 'SUELTO':
                        billing.classify_one_off(p, category)
                        summary['one_off'] += 1
                        money['one_off'] += p.amount
                        note('Pagos', r, 'ok', f'pago #{p.id}: pago suelto ({billing.ONE_OFF_CATEGORIES[category]})')
                        continue
                    contract = _contract_for(key, p.patient_id, contracts)
                    inst_id = _installment_id(contract, number)
                    made = billing.allocate_existing(
                        p, contract_id=contract.id, installment_id=inst_id, import_id=import_id, user_id=user_id
                    )
                    p.import_id = p.import_id or import_id
                    touched.add(p.patient_id)
                    applied = round(sum(a.amount for a in made), 2)
                    summary['applied'] += 1
                    money['applied'] += applied
                    extra = p.amount - applied
                    note(
                        'Pagos',
                        r,
                        'ok',
                        f'pago #{p.id}: S/ {applied:.2f} a {len(made)} {"cuota" if len(made) == 1 else "cuotas"} del contrato {key or "#" + str(contract.id)}'
                        + (f'; S/ {extra:.2f} quedan a favor' if extra > billing.EPS else ''),
                    )
                    continue

                # Pago nuevo
                u = _patient(rec, by_code)
                when_ = _cell_date(rec.get('Fecha'))
                amount = _num(rec.get('Monto'), allow_empty=False)
                if not when_:
                    raise ValueError('falta la fecha')
                ref = str(rec.get('Referencia') or '').strip() or None
                dup = Payment.query.filter(
                    Payment.patient_id == u.id,
                    Payment.amount == amount,
                    db.func.date(Payment.date) == when_,
                    Payment.is_active.isnot(False),
                    (Payment.reference == ref) if ref else db.true(),
                ).first()
                if dup:
                    summary['duplicates'] += 1
                    note(
                        'Pagos',
                        r,
                        'skip',
                        f'posible duplicado del pago #{dup.id} (mismo paciente, fecha y monto): no se crea',
                    )
                    continue
                kw = {
                    'date_': datetime.combine(when_, datetime.min.time()).replace(hour=12),
                    'reference': ref,
                    'discount': _num(rec.get('Descuento')) or 0,
                    'notes': (str(rec.get('Comentario') or rec.get('Notas') or '').strip() or None),
                    'source': 'import',
                    'import_id': import_id,
                    'user_id': user_id,
                    'commit': False,
                }
                method = str(rec.get('Método') or '').strip() or 'Otro'
                if dest == 'SUELTO':
                    billing.apply_patient_payment(u.id, amount, method, kind='one_off', category=category, **kw)
                    summary['new_one_off'] += 1
                    money['one_off'] += amount
                    note(
                        'Pagos',
                        r,
                        'ok',
                        f'pago nuevo de S/ {amount:.2f}: suelto ({billing.ONE_OFF_CATEGORIES.get(category, category)})',
                    )
                else:
                    contract = _contract_for(key, u.id, contracts)
                    p, made, left = billing.apply_patient_payment(
                        u.id,
                        amount,
                        method,
                        contract_id=contract.id,
                        installment_id=_installment_id(contract, number),
                        **kw,
                    )
                    touched.add(u.id)
                    summary['new_applied'] += 1
                    money['applied'] += amount - left
                    note(
                        'Pagos',
                        r,
                        'ok',
                        f'pago nuevo de S/ {amount:.2f} a {len(made)} cuota(s)'
                        + (f'; S/ {left:.2f} a favor' if left > billing.EPS else ''),
                    )
            except (ValueError, billing.BillingError) as e:
                note('Pagos', r, 'error', str(e))

        for pid in touched:
            billing.apply_credit(pid, source='credit', import_id=import_id, user_id=user_id)
        credits = {pid: billing.credit(pid) for pid in touched}

        # 3) Cómo queda cada contrato nuevo: si la primera cuota quedó muy atrás, se ve aquí como deuda vencida.
        db.session.flush()
        today = date.today()
        for key, c in contracts.items():
            if key.startswith('#'):
                continue
            insts = [i for i in c.installments if i.is_active is not False]
            overdue = [i for i in insts if i.due_date < today and billing._remaining(i) > billing.EPS]
            owed = round(sum(billing._remaining(i) for i in overdue), 2)
            paid = sum(1 for i in insts if i.status == 'paid')
            msg = f'{key}: {paid} de {len(insts)} cuotas pagadas'
            msg += f'; quedarían S/ {owed:.2f} vencidos en {len(overdue)} cuota(s)' if owed else '; sin deuda vencida'
            if credits.get(c.patient_id):
                msg += f'; saldo a favor S/ {credits[c.patient_id]:.2f}'
            note('Contratos', 0, 'warn' if owed else 'info', msg)
            summary['overdue_after'] += owed

        result_summary = {
            **summary,
            'amount_applied': round(money['applied'], 2),
            'amount_one_off': round(money['one_off'], 2),
            'patients': len(touched),
            'credit_total': round(sum(credits.values()), 2),
            'overdue_after': round(summary.get('overdue_after', 0), 2),
            'errors': errors,
        }
        if apply and not errors:
            batch.summary = json.dumps(result_summary)
            db.session.commit()
            return {
                'ok': True,
                'applied': True,
                'import_id': batch.id,
                'sha256': sha,
                'summary': result_summary,
                'rows': out_rows,
            }
        db.session.rollback()
        return {'ok': errors == 0, 'applied': False, 'sha256': sha, 'summary': result_summary, 'rows': out_rows}
    except Exception:
        db.session.rollback()
        raise


def _contract_for(key, patient_id, contracts):
    if key:
        c = contracts.get(key)
        if not c:
            raise ValueError(f'el contrato «{key}» no está en la hoja Contratos (o tiene errores)')
        if c.patient_id != patient_id:
            raise ValueError(f'el contrato {key} es de otro paciente')
        return c
    mine = billing.active_contracts(patient_id)
    if len(mine) == 1:
        return mine[0]
    raise ValueError(
        'indica en «Contrato (clave)» a qué contrato va (el paciente tiene ' + ('varios' if mine else 'ninguno') + ')'
    )


def _installment_id(contract, number):
    if number in (None, ''):
        return None
    for inst in contract.installments:
        if inst.number == int(number):
            return inst.id
    raise ValueError(f'el contrato #{contract.id} no tiene la cuota n° {number}')


def recent_imports(limit=10):
    return [
        {
            'id': b.id,
            'file_name': b.file_name,
            'created_at': b.created_at.isoformat() + 'Z',
            'summary': json.loads(b.summary or '{}'),
            'allocations': PaymentAllocation.query.filter_by(import_id=b.id).count(),
        }
        for b in BillingImport.query.order_by(BillingImport.id.desc()).limit(limit)
    ]
