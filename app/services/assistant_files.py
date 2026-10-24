"""Lo que se sube al chat del asistente (PRP 016): fotos y documentos.

- Foto de un comprobante (Yape, transferencia…) → formulario de «registrar pago» con monto, fecha, método,
  operación y el paciente ya elegido si el nombre coincide; boleta o factura → formulario de gasto.
- Documento (PDF, Word, texto, Excel) → resumen hecho por la IA y la propuesta de guardarlo en el Drive (y después
  imprimirlo).

La lectura de imágenes usa MiMo (visión) y, si no hay cuota, el OCR de respaldo (Gemini/Groq). El texto de los
documentos se recorta antes de mandarlo al modelo: los modelos gratis tienen poco contexto.
"""

import os
import re
import uuid

from flask import current_app

from app.services import assistant_analytics as analytics
from app.services.tools_files import UPLOADS_DIR, upload_dir, upload_path

IMAGE_EXT = {'jpg', 'jpeg', 'png', 'webp'}
DOC_EXT = {'pdf', 'docx', 'txt', 'md', 'csv', 'xlsx'}
MAX_BYTES = 15 * 1024 * 1024
MAX_TEXT = 12000
METHODS = {'yape': 'Yape', 'plin': 'Transferencia', 'transferencia': 'Transferencia', 'efectivo': 'Efectivo'}


class FileError(ValueError):
    pass


def _noop(*_a, **_k):
    return None


def save_upload(storage, user):
    """Guarda el archivo subido y devuelve su ficha. El nombre lleva el id del usuario (ver `upload_path`)."""
    original = os.path.basename(storage.filename or 'archivo')
    ext = original.rsplit('.', 1)[-1].lower() if '.' in original else ''
    if ext not in IMAGE_EXT | DOC_EXT:
        raise FileError('Sube una foto (JPG, PNG, WEBP) o un documento (PDF, Word, Excel, TXT).')
    data = storage.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise FileError('El archivo pasa de 15 MB.')
    safe = re.sub(r'[^\w.\- ]+', '', original)[:80] or f'archivo.{ext}'
    token = f'{user.id}_{uuid.uuid4().hex[:10]}_{safe}'
    with open(os.path.join(upload_dir(), token), 'wb') as fh:
        fh.write(data)
    kind = 'image' if ext in IMAGE_EXT else 'document'
    return {'token': token, 'name': original, 'kind': kind, 'size': len(data), 'url': f'/uploads/mcp/{token}'}


# ── IA para texto ───────────────────────────────────────────────────────────
def ai_text(prompt, user, model=None):
    """(texto, quién) con MiMo si está; si no (o si eligieron otro modelo), con la cadena de respaldo."""
    from app.services import llm_client, opencode_agent
    from app.services import opencode_client as oc

    if model in (None, '', 'auto', 'mimo') and user is not None and opencode_agent.enabled_for(user):
        try:
            sid = oc.create_session(f'moscowle texto u{user.id}')
            try:
                res = oc.send(sid, prompt, timeout=90)
            finally:
                oc.delete_session(sid)
            if (res.get('text') or '').strip():
                return res['text'].strip(), 'MiMo (opencode)'
        except oc.OpencodeUnavailableError as exc:
            current_app.logger.info(f'asistente: MiMo no disponible para texto: {exc}')
    token = llm_client.PREFERRED_PROVIDER.set(None if model in (None, '', 'auto', 'mimo') else model)
    try:
        content, provider = llm_client.llm_chat(
            [{'role': 'user', 'content': prompt}], temperature=0.2, max_tokens=2000, phase='resume'
        )
        return (content or '').strip() or None, f'{provider} (respaldo)'
    except Exception as exc:  # noqa: BLE001 - sin modelo: el que llama sigue sin la parte redactada
        current_app.logger.warning(f'asistente: sin modelo para texto: {exc}')
        return None, None
    finally:
        llm_client.PREFERRED_PROVIDER.reset(token)


# ── documentos ──────────────────────────────────────────────────────────────
def extract_text(path):
    ext = path.rsplit('.', 1)[-1].lower()
    if ext == 'pdf':
        try:
            from pypdf import PdfReader
        except ImportError:  # sin la librería: el pdftotext del sistema (poppler), si está
            import subprocess

            cmd = ['pdftotext', '-l', '40', path, '-']  # ruta propia del servidor, sin entrada libre
            out = subprocess.run(cmd, capture_output=True, text=True, timeout=60, check=False)  # noqa: S603
            return out.stdout
        reader = PdfReader(path)
        return '\n'.join((page.extract_text() or '') for page in reader.pages[:40])
    if ext == 'docx':
        import docx

        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            parts += [' | '.join(c.text for c in row.cells) for row in table.rows]
        return '\n'.join(parts)
    if ext == 'xlsx':
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets[:3]:
            out.append(f'# Hoja {ws.title}')
            for row in ws.iter_rows(max_row=200, values_only=True):
                out.append(' | '.join('' if v is None else str(v) for v in row))
        return '\n'.join(out)
    with open(path, encoding='utf-8', errors='replace') as fh:
        return fh.read(MAX_TEXT * 2)


def _person_match(name, user):
    """Id del paciente cuyo nombre coincide con `name` (todas las palabras del más corto en el otro)."""
    from app.services.assistant_service import PERSON_FIELDS, _people

    words = set(analytics._norm(name or '').split())
    if len(words) < 2:
        return None
    hits = []
    for o in _people(PERSON_FIELDS['patient_id'], user):
        cand = set(analytics._norm(o['label'].split(' · ')[0]).split())
        small, big = (cand, words) if len(cand) <= len(words) else (words, cand)
        if len(small) >= 2 and small <= big:
            hits.append(o['value'])
    return hits[0] if len(hits) == 1 else None


def _read_image(path, user, model, step):
    from app.services import opencode_agent
    from app.services.opencode_client import OpencodeUnavailableError

    if model in (None, '', 'auto', 'mimo') and opencode_agent.enabled_for(user):
        step('Leyendo la imagen con MiMo…')
        try:
            return opencode_agent.extract_image(user, path), 'MiMo (opencode)'
        except OpencodeUnavailableError as exc:
            current_app.logger.info(f'asistente: visión de MiMo no disponible: {exc}')
    step('Leyendo la imagen con el OCR de respaldo…')
    from app.routes.mcp_routes import _ocr_voucher

    ocr = _ocr_voucher(path, path.rsplit('.', 1)[-1].lower()) or {}
    if not ocr.get('amount'):
        return {'doc_type': 'other', 'fields': {}, 'warnings': [], 'confidence': 0}, 'OCR de respaldo'
    fields = {**ocr, 'payer': ocr.get('patient_hint')}
    return {'doc_type': 'voucher', 'fields': fields, 'warnings': [], 'confidence': 0.5}, 'OCR de respaldo'


def analyze(att, note, user, step=_noop, model=None):
    """Respuesta del asistente para un archivo subido (mismo formato que `compute`)."""
    from app.services.assistant_service import _action_answer

    token = att.get('token')
    try:
        path = upload_path(token, user.id)
    except ValueError as exc:
        raise FileError(str(exc)) from exc
    name = att.get('name') or os.path.basename(path)
    ext = path.rsplit('.', 1)[-1].lower()
    is_admin = user.role == 'admin'

    if ext in IMAGE_EXT:
        intake, who = _read_image(path, user, model, step)
        f = intake.get('fields') or {}
        warn = ''.join(f'\n- ⚠️ {w}' for w in intake.get('warnings') or [])
        if intake.get('doc_type') == 'voucher':
            args = {
                'amount': f.get('amount'),
                'payment_date': f.get('date'),
                'reference': f.get('reference'),
                'method': METHODS.get(analytics._norm(f.get('method') or '')),
                'receipt_url': att.get('url') or f'/uploads/mcp/{token}',
                'patient_id': _person_match(f.get('payer'), user),
            }
            args = {k: v for k, v in args.items() if v not in (None, '')}
            payer = f.get('payer') or 'no se lee el nombre'
            intro = (
                f'Es un comprobante de pago: **S/ {f.get("amount")}** · {f.get("method") or "método sin leer"} · '
                f'{f.get("date") or "sin fecha"} · pagó: {payer}.'
                + ('' if args.get('patient_id') else ' Elige al paciente en el formulario.')
                + warn
            )
            return _action_answer('register_payment', args, user, who, intro=intro)
        if intake.get('doc_type') == 'expense_receipt' and is_admin:
            desc = ' · '.join(x for x in (f.get('vendor'), f.get('description')) if x) or 'Gasto del centro'
            args = {'description': desc[:200], 'amount': f.get('amount')}
            args = {k: v for k, v in args.items() if v not in (None, '')}
            intro = f'Es una boleta o factura: **S/ {f.get("amount") or "?"}** · {desc}.{warn}'
            return _action_answer('create_expense', args, user, who, intro=intro)
        text = 'No parece un comprobante ni una boleta.' + warn
        blocks = [{'type': 'text', 'md': text}]
        if is_admin:
            blocks += _drive_offer(token, name, user)
        return {'title': name[:80], 'blocks': blocks, 'sources': [{'tool': 'vision', 'scope': 'imagen'}], 'model': who}

    # Documento
    step('Leyendo el documento…')
    try:
        content = extract_text(path)
    except Exception as exc:  # noqa: BLE001
        current_app.logger.warning(f'asistente: no se pudo leer {name}: {exc}')
        content = ''
    content = re.sub(r'\n{3,}', '\n\n', content or '').strip()
    if not content:
        summary, who = 'No pude leer texto en este documento (¿es un escaneo?).', 'lectura del documento'
    else:
        step('Resumiendo el documento…')
        ask = (note or '').strip()
        prompt = (
            'Eres el asistente de un centro de terapias en Perú. Lee el documento y responde en español, breve:\n'
            '1) Qué es (tipo de documento) y de quién/qué paciente trata, si se ve.\n'
            '2) Resumen en 4-8 viñetas con lo importante (fechas, montos, acuerdos, indicaciones).\n'
            '3) Qué recomiendas hacer con él (guardarlo, imprimirlo, registrar algo en el sistema).\n'
            + (f'La persona además pregunta: «{ask}». Respóndelo también.\n' if ask else '')
            + f'\nDocumento «{name}»:\n"""\n{content[:MAX_TEXT]}\n"""'
        )
        summary, who = ai_text(prompt, user, model)
        if not summary:
            lines = [ln for ln in content.splitlines() if ln.strip()][:12]
            summary = 'No hay un modelo de IA disponible para resumirlo. Empieza así:\n\n' + '\n'.join(
                f'> {ln[:200]}' for ln in lines
            )
            who = 'lectura del documento'
    blocks = [{'type': 'text', 'md': summary}]
    if is_admin:
        blocks += _drive_offer(token, name, user)
    return {'title': name[:80], 'blocks': blocks, 'sources': [{'tool': 'documento', 'scope': name}], 'model': who}


def _drive_offer(token, name, user):
    from app.services.assistant_service import action_block

    return [
        {'type': 'text', 'md': '¿Lo guardo en el Drive? Después lo puedes imprimir desde aquí.'},
        action_block('drive_save_upload', {'upload': token, 'folder': UPLOADS_DIR, 'name': name}, user),
    ]
