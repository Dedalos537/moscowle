"""Libro de cobranza: todo pago se aplica a cuotas (FIFO), deja saldo a favor y el Excel importa con trazabilidad."""

import io
import uuid
from datetime import date, datetime, timedelta

import pytest
from flask_jwt_extended import create_access_token
from openpyxl import load_workbook

from app.extensions import db
from app.models import Contract, Installment, Payment, User
from app.models.billing import BillingImport, PaymentAllocation
from app.services import billing
from app.services.billing_excel import export_workbook, import_workbook
from app.services.contract_debt import contract_debt
from app.services.contract_service import ContractService

_MADE: list = []


def _user(role='jugador', **kw):
    tag = uuid.uuid4().hex[:8]
    u = User(username=f'{role[:3]}{tag}', email=f'{tag}@bill.test', password='x', role=role, **kw)
    db.session.add(u)
    db.session.commit()
    _MADE.append(u)
    return u


@pytest.fixture(autouse=True)
def clean(app):
    _MADE.clear()
    yield
    db.session.rollback()
    ids = [u.id for u in _MADE]
    if not ids:
        return
    pids = [p for (p,) in db.session.query(Payment.id).filter(Payment.patient_id.in_(ids))]
    if pids:
        PaymentAllocation.query.filter(PaymentAllocation.payment_id.in_(pids)).delete(synchronize_session=False)
    cids = [c for (c,) in db.session.query(Contract.id).filter(Contract.patient_id.in_(ids))]
    if cids:
        Installment.query.filter(Installment.contract_id.in_(cids)).update(
            {'payment_id': None}, synchronize_session=False
        )
    if pids:
        Payment.query.filter(Payment.id.in_(pids)).delete(synchronize_session=False)
    if cids:
        Installment.query.filter(Installment.contract_id.in_(cids)).delete(synchronize_session=False)
        Contract.query.filter(Contract.id.in_(cids)).delete(synchronize_session=False)
    for u in User.query.filter(User.id.in_(ids)):
        u.is_active = False
    db.session.commit()


def _plan(patient, first_due, n=3, amount=150):
    return billing.create_plan(patient.id, amount, first_due.day, first_due, installments=n)


def _admin_headers():
    admin = _user('admin')
    return {'Authorization': 'Bearer ' + create_access_token(identity=str(admin.id)), 'Accept': 'application/json'}


def test_payment_from_register_covers_oldest_installments_first(app, client):
    p = _user()
    today = date.today()
    c = _plan(p, today - timedelta(days=70), n=4)
    assert contract_debt(today, [p.id])[p.id].overdue_amount >= 300

    r = client.post(
        '/admin/payments/register',
        data={'patient_id': p.id, 'amount': '300', 'method': 'Yape', 'payment_date': today.isoformat()},
        headers=_admin_headers(),
    )
    body = r.get_json()
    assert r.status_code == 200 and body['success'] and len(body['applied']) == 2
    first, second = sorted(c.installments, key=lambda i: i.number)[:2]
    db.session.refresh(first)
    db.session.refresh(second)
    assert first.status == 'paid' and second.status == 'paid'
    debt = contract_debt(today, [p.id])[p.id]
    assert debt.overdue_amount == pytest.approx(
        max(0, 150 * sum(1 for i in c.installments if i.due_date < today) - 300)
    )


def test_overpayment_becomes_credit_and_covers_the_next_installment(app):
    p = _user()
    today = date.today()
    first_due = date(today.year, today.month, 1)
    c = billing.create_plan(p.id, 100, 1, first_due, installments=None, today=today)  # renovable
    open_now = billing.open_installments(p.id)
    pay, made, left = billing.apply_patient_payment(p.id, 100 * len(open_now) + 60, 'Efectivo')
    assert left == pytest.approx(60) and billing.credit(p.id) == pytest.approx(60)
    # Llega el mes siguiente: la nueva cuota nace cubierta en parte con el saldo a favor.
    billing.extend_renewables(today + timedelta(days=40))
    partial = [i for i in Installment.query.filter_by(contract_id=c.id) if i.status == 'partial']
    assert len(partial) == 1 and partial[0].paid_amount == pytest.approx(60)
    assert billing.credit(p.id) == pytest.approx(0)


def test_without_contract_requires_a_plan_or_a_classified_one_off(app, client):
    p = _user()
    headers = _admin_headers()
    r = client.post(
        '/admin/payments/register', data={'patient_id': p.id, 'amount': '80', 'method': 'Yape'}, headers=headers
    )
    assert r.status_code == 400 and r.get_json()['code'] == 'no_contract'
    r = client.post(
        '/admin/payments/register',
        data={'patient_id': p.id, 'amount': '80', 'method': 'Yape', 'kind': 'one_off'},
        headers=headers,
    )
    assert r.status_code == 400 and r.get_json()['code'] == 'category'
    r = client.post(
        '/admin/payments/register',
        data={'patient_id': p.id, 'amount': '80', 'method': 'Yape', 'kind': 'one_off', 'category': 'evaluacion'},
        headers=headers,
    )
    assert r.status_code == 200
    pay = db.session.get(Payment, r.get_json()['payment_id'])
    assert pay.kind == 'one_off' and pay.category == 'evaluacion' and billing.allocated(pay.id) == 0
    # Validación: error en JSON (antes era una redirección HTML).
    r = client.post(
        '/admin/payments/register', data={'patient_id': p.id, 'amount': '', 'method': 'Yape'}, headers=headers
    )
    assert r.status_code == 400 and 'monto' in r.get_json()['error']


def test_installment_payment_is_traced_and_reverts_cleanly(app):
    p = _user()
    c = _plan(p, date.today() - timedelta(days=5), n=2)
    inst = sorted(c.installments, key=lambda i: i.number)[1]
    ok, pay = ContractService().pay_installment(inst.id, 150, 'Plin')
    assert ok and PaymentAllocation.query.filter_by(payment_id=pay.id, installment_id=inst.id).count() == 1
    db.session.refresh(inst)
    assert inst.status == 'paid'
    ContractService().revert_payment(pay)
    db.session.commit()
    db.session.refresh(inst)
    assert inst.status == 'pending' and inst.paid_amount == 0 and billing.allocated(pay.id) == 0


def test_waived_month_counts_neither_as_collected_nor_as_debt(app):
    p = _user()
    c = _plan(p, date.today() - timedelta(days=40), n=2)
    inst = sorted(c.installments, key=lambda i: i.number)[0]
    ok, _ = ContractService().pay_installment(inst.id, 0, 'Acuerdo', is_free_month=True)
    db.session.refresh(inst)
    assert ok and inst.status == 'free' and inst.amount == 0 and 'S/ 150.00' in inst.payment_notes


def test_excel_round_trip_preview_apply_and_idempotent(app):
    p = _user(login_code='PCJP' + uuid.uuid4().hex[:5].upper(), payment_amount=150, payment_plan='monthly')
    when = datetime.now() - timedelta(days=45)
    old = [
        Payment(patient_id=p.id, amount=150, method='Yape', date=when, kind='unassigned'),
        Payment(patient_id=p.id, amount=150, method='Yape', date=when + timedelta(days=30), kind='unassigned'),
        Payment(patient_id=p.id, amount=60, method='Efectivo', date=when + timedelta(days=2), kind='unassigned'),
    ]
    db.session.add_all(old)
    db.session.commit()

    wb = load_workbook(io.BytesIO(export_workbook()))
    pays, cons = wb['Pagos'], wb['Contratos']
    key = f'C-{p.login_code}'
    mine = {o.id for o in old}
    # Deja solo las filas de este test (la base de tests es compartida).
    for r in range(pays.max_row, 1, -1):
        if pays.cell(r, 1).value not in mine:
            pays.delete_rows(r)
    for r in range(cons.max_row, 1, -1):
        if cons.cell(r, 1).value != key:
            cons.delete_rows(r)
    assert cons.max_row == 2 and cons.cell(2, 8).value == 150  # contrato propuesto con el monto del plan
    for r in range(2, pays.max_row + 1):
        if pays.cell(r, 7).value == 60:
            pays.cell(r, 13).value, pays.cell(r, 14).value = 'SUELTO', 'material'
        else:
            pays.cell(r, 13).value, pays.cell(r, 15).value = 'CUOTAS', key
    cons.cell(2, 9).value = 2  # dos cuotas
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    preview = import_workbook(data, file_name='t.xlsx', apply=False)
    assert preview['ok'] and not preview['applied']
    assert Contract.query.filter_by(patient_id=p.id).count() == 0  # la vista previa no toca la base
    assert preview['summary']['contracts_created'] == 1 and preview['summary']['applied'] == 2

    done = import_workbook(data, file_name='t.xlsx', apply=True)
    assert done['applied'] and done['summary']['one_off'] == 1
    c = Contract.query.filter_by(patient_id=p.id).one()
    assert all(i.status == 'paid' for i in c.installments)
    assert PaymentAllocation.query.filter_by(import_id=done['import_id']).count() == 2
    assert db.session.get(BillingImport, done['import_id']).file_sha256 == done['sha256']

    again = import_workbook(data, file_name='t.xlsx', apply=True)
    assert again['applied'] and again['summary'].get('contracts_created', 0) == 0
    assert again['summary'].get('already_applied', 0) == 2 and Contract.query.filter_by(patient_id=p.id).count() == 1


def test_voucher_is_read_only_with_mimo_and_fills_the_form(app, client, monkeypatch):
    """El voucher se lee SOLO con MiMo; si no está, se avisa (no se prueba otro modelo)."""
    from app.services import opencode_agent
    from app.services.opencode_client import OpencodeUnavailableError

    headers = _admin_headers()
    fields = {'amount': 350, 'date': '2026-10-19', 'method': 'Yape', 'reference': '00123', 'payer': 'Ana'}
    monkeypatch.setattr(
        opencode_agent,
        'extract_image',
        lambda user, path: {'doc_type': 'voucher', 'confidence': 0.9, 'fields': fields, 'warnings': []},
    )
    photo = (io.BytesIO(b'\xff\xd8fake-jpeg'), 'voucher.jpg')
    r = client.post('/admin/analyze-receipt', data={'receipt': photo}, headers=headers)
    body = r.get_json()
    assert body['source'] == 'mimo' and body['amount'] == 350 and body['method'] == 'yape'
    assert body['date'] == '2026-10-19' and body['reference'] == '00123'

    def down(user, path):
        raise OpencodeUnavailableError('quota')

    monkeypatch.setattr(opencode_agent, 'extract_image', down)
    photo = (io.BytesIO(b'\xff\xd8fake-jpeg'), 'voucher.jpg')
    body = client.post('/admin/analyze-receipt', data={'receipt': photo}, headers=headers).get_json()
    assert 'MiMo no está disponible' in body['warning'] and body.get('amount') is None


def test_contract_detail_lists_each_payment_of_an_installment_with_its_voucher(app):
    p = _user()
    c = _plan(p, date.today(), n=1, amount=500)
    inst = c.installments[0]
    billing.apply_patient_payment(p.id, 350, 'Yape', receipt_path='vouchers/a.jpg', reference='OP1')
    billing.apply_patient_payment(p.id, 100, 'Efectivo')
    detail = ContractService().get_contract_detail(c.id)
    row = next(i for i in detail['installments'] if i['id'] == inst.id)
    assert row['status'] == 'partial' and row['paid_amount'] == pytest.approx(450)
    assert [x['amount'] for x in row['payments']] == [350, 100]
    assert row['payments'][0]['receipt_image_path'] == 'vouchers/a.jpg' and row['payments'][0]['reference'] == 'OP1'
    assert row['receipt_image_path'] == 'vouchers/a.jpg'  # antes un pago parcial no mostraba su voucher
