"""Drive del administrador: seguridad, operaciones de archivos, subida por partes, compresión y vistas previas."""

import io
import uuid
import zipfile
from unittest.mock import patch

import pytest
from flask_jwt_extended import create_access_token

from app.extensions import db
from app.models import User
from app.services import drive_service as drive


def _user(role):
    tag = uuid.uuid4().hex[:8]
    u = User(username=f'{role[:3]}{tag}', email=f'{tag}@drv.test', password='x', role=role)
    db.session.add(u)
    db.session.commit()
    return u


@pytest.fixture
def drv(app, client, tmp_path):
    app.config['DRIVE_ROOT'] = str(tmp_path / 'drive')
    app.config['DRIVE_QUOTA_GB'] = 30
    drive._usage.update(bytes=None, at=0)
    headers = {'Authorization': 'Bearer ' + create_access_token(identity=str(_user('admin').id))}

    class Api:
        def get(self, url, **kw):
            return client.get('/api/drive' + url, headers=headers, **kw)

        def post(self, url, json=None):
            return client.post('/api/drive' + url, headers=headers, json=json or {})

        def put(self, url, data):
            return client.put('/api/drive' + url, headers=headers, data=data)

        def delete(self, url):
            return client.delete('/api/drive' + url, headers=headers)

        def upload(self, folder, name, content: bytes):
            init = self.post('/upload/init', {'path': folder, 'name': name, 'size': len(content)}).get_json()
            uid, step, offset = init['upload_id'], 3, 0
            while offset < len(content):
                part = content[offset : offset + step]
                r = self.put(f'/upload/{uid}?offset={offset}', part)
                assert r.status_code == 200, r.get_json()
                offset += len(part)
            return self.post(f'/upload/{uid}/complete')

    yield Api()
    app.config.pop('DRIVE_ROOT', None)
    drive._usage.update(bytes=None, at=0)


def test_only_admin_can_use_the_drive(client, app, drv):
    sup = {'Authorization': 'Bearer ' + create_access_token(identity=str(_user('supervisor').id))}
    assert client.get('/api/drive/list', headers=sup).status_code == 403
    assert client.get('/api/drive/list').status_code == 401
    assert drv.get('/list').status_code == 200


@pytest.mark.parametrize('bad', ['../', '../../etc/passwd', '.trash', '.uploads/x', '/etc'])
def test_paths_outside_the_drive_are_rejected(drv, bad):
    r = drv.get('/list', query_string={'path': bad})
    assert r.status_code in (400, 404)
    assert drv.post('/folder', {'path': bad, 'name': 'x'}).status_code in (400, 404)


def test_folder_rename_move_copy_trash_restore(drv):
    assert drv.post('/folder', {'path': '', 'name': 'Informes'}).status_code == 201
    assert drv.post('/folder', {'path': '', 'name': 'Informes'}).get_json()['name'] == 'Informes (2)'
    assert drv.upload('Informes', 'a.txt', b'hola mundo').status_code == 201
    assert drv.post('/rename', {'path': 'Informes/a.txt', 'name': 'b.txt'}).get_json()['path'] == 'Informes/b.txt'
    assert drv.post('/rename', {'path': 'Informes/b.txt', 'name': '../x'}).status_code == 400
    assert drv.post('/copy', {'paths': ['Informes/b.txt'], 'dest': 'Informes (2)'}).status_code == 200
    assert drv.post('/move', {'paths': ['Informes (2)'], 'dest': 'Informes'}).get_json()['entries'][0]['path'] == (
        'Informes/Informes (2)'
    )
    assert drv.post('/move', {'paths': ['Informes'], 'dest': 'Informes/Informes (2)'}).status_code == 400
    names = [e['name'] for e in drv.get('/list', query_string={'path': 'Informes'}).get_json()['entries']]
    assert names == ['Informes (2)', 'b.txt']  # carpetas primero
    assert drv.post('/delete', {'paths': ['Informes/b.txt']}).get_json()['trashed'] == 1
    trash = drv.get('/trash').get_json()['entries']
    assert trash[0]['original'] == 'Informes/b.txt'
    assert drv.post('/trash/restore', {'ids': [trash[0]['id']]}).get_json()['restored'] == 1
    assert drv.get('/file', query_string={'path': 'Informes/b.txt'}).data == b'hola mundo'


def test_upload_is_ordered_and_respects_quota(app, drv):
    init = drv.post('/upload/init', {'path': '', 'name': 'x.bin', 'size': 10}).get_json()
    assert drv.put(f'/upload/{init["upload_id"]}?offset=5', b'abc').status_code == 409  # fuera de orden
    assert drv.post(f'/upload/{init["upload_id"]}/complete').status_code == 409  # incompleta
    app.config['DRIVE_QUOTA_GB'] = 1 / 1024**3 * 100  # 100 bytes
    drive._usage.update(bytes=None, at=0)
    assert drv.post('/upload/init', {'path': '', 'name': 'grande.bin', 'size': 500}).status_code == 507


def test_compress_and_extract_roundtrip(drv):
    drv.post('/folder', {'path': '', 'name': 'Fotos'})
    drv.upload('Fotos', 'uno.txt', b'1')
    drv.upload('Fotos', 'dos.txt', b'22')
    z = drv.post('/compress', {'paths': ['Fotos'], 'dest': '', 'name': 'Fotos'}).get_json()
    assert z['name'] == 'Fotos.zip'
    out = drv.post('/extract', {'path': 'Fotos.zip'}).get_json()
    assert out['name'] == 'Fotos (2)'
    inner = [e['name'] for e in drv.get('/list', query_string={'path': 'Fotos (2)/Fotos'}).get_json()['entries']]
    assert inner == ['dos.txt', 'uno.txt']


def test_zip_slip_is_rejected(drv):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as zf:
        zf.writestr('../../escape.txt', 'x')
    drv.upload('', 'malo.zip', buf.getvalue())
    r = drv.post('/extract', {'path': 'malo.zip'})
    assert r.status_code == 400 and 'peligrosa' in r.get_json()['error']


def test_previews_for_text_docx_and_pptx(drv):
    import docx

    drv.upload('', 'nota.txt', 'línea 1\nlínea 2'.encode())
    assert drv.get('/preview', query_string={'path': 'nota.txt'}).get_json()['text'].startswith('línea 1')

    d = docx.Document()
    d.add_heading('Informe mensual', level=1)
    d.add_paragraph('Texto <script>alert(1)</script>')
    buf = io.BytesIO()
    d.save(buf)
    drv.upload('', 'informe.docx', buf.getvalue())
    with patch('app.services.drive_preview.soffice_bin', return_value=None):
        prev = drv.get('/preview', query_string={'path': 'informe.docx'}).get_json()
    assert prev['mode'] == 'html' and '<h1>Informe mensual</h1>' in prev['html']
    assert '<script>' not in prev['html'] and '&lt;script&gt;' in prev['html']

    ns = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    slide = f'<p:sld {ns}><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>Bienvenida</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>'
    pbuf = io.BytesIO()
    with zipfile.ZipFile(pbuf, 'w') as zf:
        zf.writestr('ppt/slides/slide1.xml', slide)
    drv.upload('', 'charla.pptx', pbuf.getvalue())
    with patch('app.services.drive_preview.soffice_bin', return_value=None):
        slides = drv.get('/preview', query_string={'path': 'charla.pptx'}).get_json()
    assert slides['mode'] == 'slides' and slides['slides'][0]['title'] == 'Bienvenida'


def test_image_thumbnail(drv):
    from PIL import Image

    buf = io.BytesIO()
    Image.new('RGB', (800, 600), 'red').save(buf, 'PNG')
    drv.upload('', 'foto.png', buf.getvalue())
    r = drv.get('/thumb', query_string={'path': 'foto.png'})
    assert r.status_code == 200 and r.mimetype == 'image/webp'


def test_printing_without_cups_explains_and_validates(drv):
    with (
        patch('app.services.print_service._bin', return_value=None),
        patch('app.services.print_service._ipp_server', return_value=None),
    ):
        status = drv.get('/printers').get_json()
        assert status['available'] is False and 'cups' in status['help'].lower()
        drv.upload('', 'doc.pdf', b'%PDF-1.4')
        r = drv.post('/print', {'path': 'doc.pdf', 'printer': 'Oficina', 'copies': 1})
        assert r.status_code == 503
    assert drv.post('/print', {'path': 'doc.pdf', 'printer': 'x; rm -rf /', 'copies': 1}).status_code == 400


def test_print_job_is_sent_with_safe_arguments(drv):
    drv.upload('', 'doc.pdf', b'%PDF-1.4')
    calls = []

    def fake_run(args, timeout=20):
        calls.append(args)
        out = 'printer Oficina is idle.\n' if args[-1] == '-p' else ''
        if args[0].endswith('lp'):
            out = 'request id is Oficina-42 (1 file(s))'
        return type('R', (), {'returncode': 0, 'stdout': out, 'stderr': ''})()

    with (
        patch('app.services.print_service._bin', side_effect=lambda n: f'/usr/bin/{n}'),
        patch('app.services.print_service._run', side_effect=fake_run),
    ):
        r = drv.post(
            '/print', {'path': 'doc.pdf', 'printer': 'Oficina', 'copies': 2, 'duplex': True, 'page_range': '1-3'}
        )
    job = r.get_json()
    assert r.status_code == 201 and job['status'] == 'sent' and job['cups_job'] == 'Oficina-42'
    lp = next(a for a in calls if a[0].endswith('/lp'))
    assert lp[:4] == ['/usr/bin/lp', '-d', 'Oficina', '-n'] and 'sides=two-sided-long-edge' in lp and '--' in lp
    assert 'print-quality=4' in lp  # normal por defecto
    bad = drv.post('/print', {'path': 'doc.pdf', 'printer': 'Oficina', 'quality': 'ultra'})
    assert bad.status_code == 400


class _FakeCups:
    """CUPS de mentira que habla IPP, como el del host visto desde VM-APP.

    Responde dentro del proceso (reemplaza `urlopen`): un servidor HTTP real en un hilo se cuelga con eventlet en CI.
    """

    address = 'cups.test:631'

    def __init__(self):
        self.requests = []

    def urlopen(self, req, timeout=None):
        import struct
        import urllib.error

        from app.services import ipp_client as ipp

        if not req.full_url.startswith(f'http://{self.address}/'):
            raise urllib.error.URLError('connection refused')
        body = req.data
        op = struct.unpack('>H', body[2:4])[0]
        _, groups = ipp.parse(body)
        self.requests.append((req.full_url.split(self.address, 1)[1], op, groups, body))
        out = struct.pack('>BBHI', 2, 0, 0, 1) + b'\x01' + ipp._attr(0x47, 'attributes-charset', 'utf-8')
        if op == ipp.CUPS_GET_PRINTERS:
            out += b'\x04' + ipp._attr(0x42, 'printer-name', 'Epson_L4150') + ipp._attr(0x23, 'printer-state', 3)
            out += b'\x04' + ipp._attr(0x42, 'printer-name', 'Oficina') + ipp._attr(0x23, 'printer-state', 5)
        elif op == ipp.CUPS_GET_DEFAULT:
            out += b'\x04' + ipp._attr(0x42, 'printer-name', 'Epson_L4150')
        elif op == ipp.PRINT_JOB:
            out += b'\x02' + ipp._attr(0x21, 'job-id', 77)
        return _Resp(out + b'\x03')


class _Resp:
    def __init__(self, data):
        self.data = data

    def read(self):
        return self.data

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def test_printing_through_host_cups_over_ipp(drv):
    cups = _FakeCups()
    with (
        patch('app.services.print_service._bin', return_value=None),
        patch('app.services.print_service._ipp_server', return_value=cups.address),
        patch('app.services.ipp_client.urllib.request.urlopen', side_effect=cups.urlopen),
    ):
        status = drv.get('/printers').get_json()
        assert status['available'] is True and status['default'] == 'Epson_L4150'
        assert [(p['name'], p['state']) for p in status['printers']] == [
            ('Epson_L4150', 'idle'),
            ('Oficina', 'disabled'),
        ]
        drv.upload('', 'doc.pdf', b'%PDF-1.4 hola')
        r = drv.post(
            '/print',
            {
                'path': 'doc.pdf',
                'printer': 'Epson_L4150',
                'copies': 2,
                'duplex': True,
                'color': False,
                'page_range': '1-3,5',
                'quality': 'high',
            },
        )
        job = r.get_json()
        assert r.status_code == 201 and job['status'] == 'sent' and job['cups_job'] == 'Epson_L4150-77'
        path, _, groups, body = cups.requests[-1]
        attrs = {k: v for g in groups for k, v in g['attrs'].items()}
        assert path == '/printers/Epson_L4150' and body.endswith(b'%PDF-1.4 hola')
        assert attrs['copies'] == [2] and attrs['sides'] == ['two-sided-long-edge']
        assert attrs['print-color-mode'] == ['monochrome'] and len(attrs['page-ranges']) == 2
        assert attrs['print-quality'] == [5] and job['quality'] == 'high'
        assert drv.post(f'/print/jobs/{job["id"]}/cancel').get_json()['status'] == 'cancelled'
        assert cups.requests[-1][1] == 0x0008


def test_unreachable_host_cups_explains_how_to_share(drv):
    cups = _FakeCups()
    with (
        patch('app.services.print_service._bin', return_value=None),
        patch('app.services.print_service._ipp_server', return_value='otro.test:631'),
        patch('app.services.ipp_client.urllib.request.urlopen', side_effect=cups.urlopen),
    ):
        status = drv.get('/printers').get_json()
    assert status['available'] is False and 'cupsctl' in status['help']


# ── escanear al drive ────────────────────────────────────────────────────────
def _jpeg(color):
    from PIL import Image

    buf = io.BytesIO()
    Image.new('RGB', (60, 80), color).save(buf, 'JPEG')
    return buf.getvalue()


@pytest.fixture
def scanner(monkeypatch):
    import json

    from app.services import scan_service

    colors = iter(['red', 'green', 'blue', 'white'])

    def fake_call(method, path, body=None, timeout=30):
        if path == '/status':
            devs = [{'name': 'airscan:w1:EPSON L4150 Series', 'label': 'WSD EPSON L4150 Series'}]
            return json.dumps({'devices': devs}).encode(), 'application/json'
        return _jpeg(next(colors)), 'image/jpeg'

    monkeypatch.setenv('SCAN_TOKEN', 't')
    monkeypatch.setenv('SCAN_SERVER_URL', 'http://192.168.122.1:8631')
    monkeypatch.setattr(scan_service, '_call', fake_call)


def test_scanner_status_lists_the_epson(drv, scanner):
    st = drv.get('/scanner').get_json()
    assert st['available'] and st['devices'][0]['name'].startswith('airscan:')


def test_without_scan_service_the_drive_explains_how_to_install_it(drv, monkeypatch):
    monkeypatch.delenv('SCAN_TOKEN', raising=False)
    st = drv.get('/scanner').get_json()
    assert st['available'] is False and 'host' in st['help']


def test_scanned_pages_become_one_pdf_in_the_chosen_folder(drv, scanner):
    drv.post('/folder', {'path': '', 'name': 'Contratos'})
    sid = drv.post('/scan/page', {'resolution': 200, 'mode': 'Color'}).get_json()['scan_id']
    drv.post('/scan/page', {'scan_id': sid})
    pages = drv.post('/scan/page', {'scan_id': sid}).get_json()['pages']
    assert [p['n'] for p in pages] == [1, 2, 3]
    assert drv.get(f'/scan/{sid}/page/2').mimetype == 'image/jpeg'
    entry = drv.post(f'/scan/{sid}/finish', {'path': 'Contratos', 'name': 'Contrato Jorge', 'format': 'pdf'})
    assert entry.status_code == 201 and entry.get_json()['name'] == 'Contrato Jorge.pdf'
    names = [e['name'] for e in drv.get('/list?path=Contratos').get_json()['entries']]
    assert names == ['Contrato Jorge.pdf']
    assert drv.get('/file?path=Contratos/Contrato%20Jorge.pdf').data.startswith(b'%PDF')
    assert drv.get(f'/scan/{sid}/page/1').status_code == 404  # el borrador se borra al guardar


def test_a_page_can_be_removed_and_a_single_page_saved_as_jpg(drv, scanner):
    sid = drv.post('/scan/page', {}).get_json()['scan_id']
    drv.post('/scan/page', {'scan_id': sid})
    left = drv.delete(f'/scan/{sid}/page/1').get_json()['pages']
    assert [p['n'] for p in left] == [2]
    entry = drv.post(f'/scan/{sid}/finish', {'path': '', 'name': 'foto', 'format': 'jpg'}).get_json()
    assert entry['name'] == 'foto.jpg'


def test_invalid_scan_options_and_ids_are_rejected(drv, scanner):
    assert drv.post('/scan/page', {'resolution': 1200}).status_code == 400
    assert drv.post('/scan/0123456789abcdef/finish', {}).status_code == 404
    assert drv.get('/scan/not-an-id/page/1').status_code == 404


def test_reconnect_printer_asks_the_host_service_and_reports_the_new_ip(drv, monkeypatch):
    import json

    from app.services import scan_service

    sent = []

    def fake_call(method, path, body=None, timeout=30):
        sent.append((method, path))
        payload = {
            'name': 'Epson_L4150',
            'uri': 'socket://192.168.1.233:9100',
            'ip': '192.168.1.233',
            'reachable': True,
            'changed': True,
            'message': 'Impresora reconectada en 192.168.1.233.',
        }
        return json.dumps(payload).encode(), 'application/json'

    monkeypatch.setenv('SCAN_TOKEN', 't')
    monkeypatch.setattr(scan_service, '_call', fake_call)
    r = drv.post('/printers/repair')
    assert r.status_code == 200 and r.get_json()['ip'] == '192.168.1.233'
    assert sent == [('POST', '/printer/repair')]  # acción fija: no viaja ningún comando
